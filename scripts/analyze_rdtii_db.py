#!/usr/bin/env python3
"""Analyze RDTII Round 1/2 database Excel files."""
from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import openpyxl

BASE = Path("Hackthon Knowledge Products/Database")
FILES = [
    BASE / "ESCAP-RDTII-2.1_ Round 1 Database.xlsx",
    BASE / "ESCAP-RDTII-2.1_ Round 2 Database.xlsx",
]


def truncate(val: Any, n: int = 200) -> str:
    if val is None:
        return ""
    s = str(val).replace("\n", " ").strip()
    if len(s) > n:
        return s[: n - 3] + "..."
    return s


def find_header_row(rows: list[tuple[Any, ...]], min_nonempty: int = 3) -> int:
    best_idx = 0
    best_score = 0
    for i, row in enumerate(rows[:20]):
        nonempty = sum(1 for c in row if c is not None and str(c).strip())
        textish = sum(
            1
            for c in row
            if c is not None and isinstance(c, str) and len(str(c).strip()) > 1
        )
        score = nonempty + textish
        if score > best_score and nonempty >= min_nonempty:
            best_score = score
            best_idx = i
    return best_idx


def read_sheet(ws: openpyxl.worksheet.worksheet.Worksheet, max_rows: int | None = None) -> list[tuple[Any, ...]]:
    rows: list[tuple[Any, ...]] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        rows.append(tuple(row))
        if max_rows and i + 1 >= max_rows:
            break
    return rows


def analyze_jurisdiction_sheet(rows: list[tuple[Any, ...]], sheet_name: str) -> dict[str, Any]:
    header_idx = find_header_row(rows)
    headers = [truncate(c, 120) if c is not None else "" for c in rows[header_idx]]
    # trim trailing empty headers
    while headers and headers[-1] == "":
        headers.pop()
    data_rows = rows[header_idx + 1 :]
    records = []
    for row in data_rows:
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        rec = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            val = row[i] if i < len(row) else None
            rec[h] = val
        if rec:
            records.append(rec)

    # Detect indicator id column
    id_col = None
    for cand in ("Indicator ID", "Indicator", "ID", "Indicator No.", "No."):
        if cand in headers:
            id_col = cand
            break
    if id_col is None:
        for h in headers:
            if re.search(r"indicator", h, re.I):
                id_col = h
                break

    indicators = []
    for rec in records:
        ind_id = rec.get(id_col) if id_col else None
        indicators.append(
            {
                "indicator_id": truncate(ind_id, 50),
                "fields": {k: truncate(v, 300) for k, v in rec.items()},
            }
        )

    # Sample unique values for scoring-ish columns
    scoring_keywords = re.compile(
        r"score|point|weight|value|rating|level|tier|binary|yes.?no|0.?1|scale|assessment|result",
        re.I,
    )
    scoring_cols = [h for h in headers if scoring_keywords.search(h)]
    col_uniques: dict[str, list[str]] = {}
    for h in scoring_cols:
        vals = sorted({truncate(rec.get(h), 80) for rec in records if rec.get(h) is not None})
        col_uniques[h] = vals[:30]
        if len(vals) > 30:
            col_uniques[h].append(f"... +{len(vals)-30} more")

    return {
        "sheet_name": sheet_name,
        "header_row_index_1based": header_idx + 1,
        "row_count_data": len(records),
        "columns": headers,
        "column_count": len(headers),
        "indicator_id_column": id_col,
        "scoring_columns": scoring_cols,
        "scoring_column_unique_values": col_uniques,
        "indicators": indicators,
        "sample_first_3": indicators[:3],
    }


def analyze_methodology(rows: list[tuple[Any, ...]]) -> dict[str, Any]:
    """Parse methodology sheet - often free-form."""
    non_empty = []
    for i, row in enumerate(rows):
        cells = [truncate(c, 500) for c in row if c is not None and str(c).strip()]
        if cells:
            non_empty.append({"row": i + 1, "cells": cells})

    # Try to find structured sections
    sections: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    for item in non_empty:
        cells = item["cells"]
        # Heuristic: single cell in first column = section header
        if len(cells) == 1 and len(cells[0]) < 120:
            if current:
                sections.append(current)
            current = {"title": cells[0], "content": []}
        elif current:
            current["content"].append(cells)
        else:
            if not sections:
                sections.append({"title": "Preamble", "content": [cells]})
            else:
                sections[-1]["content"].append(cells)
    if current:
        sections.append(current)

    return {
        "non_empty_row_count": len(non_empty),
        "sections": sections[:50],
        "all_non_empty_preview": non_empty[:80],
    }


def analyze_consolidated(rows: list[tuple[Any, ...]]) -> dict[str, Any]:
    header_idx = find_header_row(rows)
    headers = [truncate(c, 120) if c is not None else "" for c in rows[header_idx]]
    while headers and headers[-1] == "":
        headers.pop()
    data_rows = rows[header_idx + 1 :]
    records = []
    for row in data_rows:
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        rec = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers)) if headers[i]}
        records.append(rec)

    # jurisdiction columns often country names
    juris_cols = [h for h in headers if h and h not in ("Indicator ID", "Indicator", "Pillar", "Sub-pillar")]
    return {
        "header_row_index_1based": header_idx + 1,
        "columns": headers,
        "row_count": len(records),
        "records_preview": [
            {k: truncate(v, 150) for k, v in r.items()} for r in records[:5]
        ],
        "all_records_summary": [
            {k: truncate(v, 100) for k, v in r.items()} for r in records
        ],
    }


def analyze_file(path: Path) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    result: dict[str, Any] = {
        "file": path.name,
        "sheet_names": wb.sheetnames,
        "sheets": {},
    }
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = read_sheet(ws)
        if "methodology" in sn.lower():
            result["sheets"][sn] = {"type": "methodology", **analyze_methodology(rows)}
        elif "consolidated" in sn.lower():
            result["sheets"][sn] = {"type": "consolidated", **analyze_consolidated(rows)}
        else:
            result["sheets"][sn] = {"type": "jurisdiction", **analyze_jurisdiction_sheet(rows, sn)}
    wb.close()
    return result


def main() -> None:
    all_results = {}
    for fp in FILES:
        print(f"Analyzing {fp}...", flush=True)
        all_results[fp.name] = analyze_file(fp)

    out = Path("scripts/rdtii_db_analysis.json")
    out.write_text(json.dumps(all_results, indent=2, default=str), encoding="utf-8")
    print(f"Wrote {out}")

    # Print human summary
    for fname, data in all_results.items():
        print("\n" + "=" * 100)
        print(f"FILE: {fname}")
        print("=" * 100)
        print(f"Sheets: {data['sheet_names']}")
        for sn, sdata in data["sheets"].items():
            print(f"\n  [{sn}] type={sdata['type']}")
            if sdata["type"] == "jurisdiction":
                print(f"    rows={sdata['row_count_data']} cols={sdata['column_count']}")
                print(f"    columns: {sdata['columns']}")
                print(f"    indicator_id_col: {sdata['indicator_id_column']}")
                print(f"    scoring_cols: {sdata['scoring_columns']}")
            elif sdata["type"] == "consolidated":
                print(f"    rows={sdata['row_count']} cols={len(sdata['columns'])}")
                print(f"    columns: {sdata['columns']}")
            elif sdata["type"] == "methodology":
                print(f"    non_empty_rows={sdata['non_empty_row_count']}")
                print(f"    sections: {[s['title'] for s in sdata['sections'][:20]]}")


if __name__ == "__main__":
    main()
